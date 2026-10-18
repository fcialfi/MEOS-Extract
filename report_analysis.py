#!/usr/bin/env python3
"""Multi-orbit SNR analysis of the Excel files produced by MEOS-Extract.

Groups passes by ground-track repeat cycle and writes, per group, the
SNR/azimuth/elevation of every orbit, ascending/descending SNR vs
elevation charts and an az/el-matched SNR table. Usable from the command
line or through :func:`run_analysis` (the GUI calls it after extraction).
"""
import argparse
import logging
import math
import re
from pathlib import Path
from typing import Dict, List, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

logger = logging.getLogger(__name__)

DEFAULT_OUTPUT_NAME = "meos_processed.xlsx"


def _open_excel(path: Path) -> pd.ExcelFile:
    try:
        return pd.ExcelFile(path)
    except Exception as exc:  # pragma: no cover - fallback for invalid engine detection
        if "io.excel.zip.reader" in str(exc):
            try:
                return pd.ExcelFile(path, engine="openpyxl")
            except Exception as inner:
                raise ValueError(
                    f"{path} appears to be a ZIP archive rather than a supported "
                    "Excel workbook. Ensure the file is a valid .xlsx/.xlsm workbook."
                ) from inner
        raise


SNR_COL = "5_10_signal_noise_ratio"
AZ_COL = "6_1_azimuth"
EL_COL = "6_2_elevation"
TIME_COL = "time_iso_utc"

# Column patterns used to find each series in a MEOS Extract workbook. Older
# reports have separate "6_1_azimuth"/"6_2_elevation" sheets, newer ones a
# single antenna sheet with "<section>_azimuth"/"<section>_elevation" columns.
# Section numbers change between report versions, so match by suffix.
SERIES_PATTERNS = {
    SNR_COL: re.compile(r"(signal_?noise_?ratio|(^|_)snr)$", re.IGNORECASE),
    AZ_COL: re.compile(r"azimuth$", re.IGNORECASE),
    EL_COL: re.compile(r"elevation$", re.IGNORECASE),
}


def _sheet_headers(path: Path) -> Dict[str, List[str]]:
    """First-row headers of every worksheet, read once per file."""
    workbook = load_workbook(path, read_only=True)
    try:
        headers = {}
        for ws in workbook.worksheets:
            first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers[ws.title] = [str(v) for v in first if v is not None]
        return headers
    finally:
        workbook.close()


def _find_series(path: Path, headers: Dict[str, List[str]], key: str) -> Tuple[str, str]:
    """Return (sheet, column) holding the series *key* (SNR, azimuth or elevation)."""
    if key in headers.get(key, []):
        return key, key
    pattern = SERIES_PATTERNS[key]
    matches = []
    for sheet, cols in headers.items():
        if sheet == "__meta__" or sheet.endswith(("_ticks", "_tick")):
            continue
        if TIME_COL not in cols:
            continue
        for col in cols:
            if pattern.search(col):
                matches.append((sheet, col))
    if not matches:
        raise ValueError(
            f"No '{key}' series found in {path.name}. "
            f"Available worksheets: {', '.join(headers) or 'None'}."
        )
    if len(matches) > 1:
        logger.warning("Several candidates for '%s' in %s: %s; using %s.", key, path.name, matches, matches[0])
    return matches[0]


def _load_series(path: Path, excel: pd.ExcelFile, sheet: str, column: str, key: str) -> pd.DataFrame:
    df = pd.read_excel(excel, sheet_name=sheet)
    if df.empty or column not in df.columns or TIME_COL not in df.columns:
        logger.warning("Sheet '%s' in %s has no '%s' data; skipping data.", sheet, path.name, column)
        return pd.DataFrame(columns=[TIME_COL, key])
    data = df[[TIME_COL, column]].rename(columns={column: key})
    data[TIME_COL] = pd.to_datetime(data[TIME_COL], errors="coerce", utc=True)
    data[key] = pd.to_numeric(data[key], errors="coerce")
    data = data.dropna().copy()
    data[TIME_COL] = data[TIME_COL].dt.floor("s")
    data = data.sort_values(TIME_COL)
    data = data.drop_duplicates(subset=[TIME_COL], keep="first")
    return data.reset_index(drop=True)


def _process_file(path: Path, match_tolerance_s: float = 2.0) -> pd.DataFrame:
    excel = _open_excel(path)
    headers = _sheet_headers(path)
    series = {key: _load_series(path, excel, *_find_series(path, headers, key), key) for key in SERIES_PATTERNS}

    # Antenna and SNR curves are sampled at different instants: pair each SNR
    # sample with the nearest azimuth/elevation sample instead of requiring
    # identical timestamps (which silently dropped rows).
    merged = series[SNR_COL]
    for key in (AZ_COL, EL_COL):
        if merged.empty or series[key].empty:
            merged = merged.iloc[0:0].assign(**{key: pd.Series(dtype=float)})
            continue
        merged = pd.merge_asof(
            merged, series[key], on=TIME_COL, direction="nearest",
            tolerance=pd.Timedelta(seconds=match_tolerance_s),
        )
    merged = merged.dropna(subset=[SNR_COL, AZ_COL, EL_COL])
    merged = merged[[TIME_COL, SNR_COL, AZ_COL, EL_COL]].copy()
    merged["orbit"] = _extract_orbit(path)
    merged = merged.sort_values(TIME_COL).reset_index(drop=True)
    return _trim_final_high_elevation(merged)


def _trim_final_high_elevation(
    frame: pd.DataFrame, threshold: float = 5.0
) -> pd.DataFrame:
    """Drop trailing samples after the pass (antenna moving back to its park position).

    The frame is cut after the last sample at or below *threshold*. This is
    done only when that sample lies in the second half of the pass: a pass
    that ends above the threshold is kept whole instead of being truncated
    to its first few samples.
    """
    if frame.empty:
        return frame
    elevation = pd.to_numeric(frame[EL_COL], errors="coerce").reset_index(drop=True)
    low = elevation.index[elevation <= threshold]
    if len(low) == 0:
        return frame.iloc[0:0].copy()
    last_low = int(low[-1])
    if last_low < len(elevation) // 2:
        logger.warning(
            "Orbit %s ends above %s deg elevation; trailing samples not trimmed.",
            frame["orbit"].iloc[0], threshold,
        )
        return frame
    return frame.iloc[: last_low + 1].copy()


def _extract_start_time(path: Path) -> pd.Timestamp:
    excel = _open_excel(path)
    if "__meta__" not in excel.sheet_names:
        available = ", ".join(excel.sheet_names)
        raise ValueError(
            f"Worksheet named '__meta__' not found in {path}. "
            f"Available worksheets: {available or 'None'}."
        )
    meta = pd.read_excel(excel, sheet_name="__meta__")
    if meta.empty:
        raise ValueError(f"Sheet '__meta__' in {path} is empty.")
    start_time = None
    if "start_time_utc" in meta.columns:
        values = meta["start_time_utc"].dropna()
        if not values.empty:
            start_time = values.iloc[0]
    if start_time is None:
        lower_cols = [str(col).strip().lower() for col in meta.columns]
        for _, row in meta.iterrows():
            for idx, col in enumerate(meta.columns):
                cell = row[col]
                if isinstance(cell, str) and cell.strip().lower() == "start_time_utc":
                    next_value = None
                    if idx + 1 < len(meta.columns):
                        next_value = row[meta.columns[idx + 1]]
                    elif "value" in lower_cols:
                        value_idx = lower_cols.index("value")
                        next_value = row[meta.columns[value_idx]]
                    start_time = next_value
                    break
            if start_time is not None:
                break
    start_time = pd.to_datetime(start_time, errors="coerce", utc=True)
    if pd.isna(start_time):
        raise ValueError(f"Unable to parse start_time_utc from {path}.")
    return start_time


def _extract_orbit(path: Path) -> int:
    match = re.search(r"orbit[_-]?(\d+)", path.stem, flags=re.IGNORECASE)
    if not match:
        raise ValueError(
            f"Unable to determine orbit from filename '{path.name}'. "
            "Expected pattern like 'orbit_2648'."
        )
    return int(match.group(1))


def _bin_series(series: pd.Series, tolerance: float) -> pd.Series:
    if tolerance <= 0:
        return series
    return (series / tolerance).round() * tolerance


def _max_center_elevation(series: pd.Series, center_fraction: float = 0.5) -> float:
    if isinstance(series, pd.DataFrame):
        series = pd.Series(series.to_numpy().ravel())
    values = pd.to_numeric(series, errors="coerce").dropna().reset_index(drop=True)
    if values.empty:
        return float("nan")
    if not (0 < center_fraction <= 1):
        center_fraction = 0.5
    window = max(1, int(round(len(values) * center_fraction)))
    start = max(0, (len(values) - window) // 2)
    end = start + window
    return float(values.iloc[start:end].max())


def _signed_elevation(series: pd.Series) -> pd.Series:
    numeric = pd.to_numeric(series, errors="coerce")
    if numeric.dropna().empty:
        return series.copy()
    max_index = numeric.idxmax()
    signed = numeric.copy()
    before_mask = signed.index < max_index
    signed[before_mask] = -signed[before_mask].abs()
    signed[~before_mask] = signed[~before_mask].abs()
    return signed


def _build_correlation_table(
    frames: List[pd.DataFrame], azimuth_tolerance: float, elevation_tolerance: float
) -> pd.DataFrame:
    aligned_frames: List[pd.DataFrame] = []
    for frame in frames:
        orbit = int(frame["orbit"].iloc[0])
        aligned = frame[
            ["6_1_azimuth", "6_2_elevation", "5_10_signal_noise_ratio"]
        ].copy()
        aligned["azimuth_bin"] = _bin_series(aligned["6_1_azimuth"], azimuth_tolerance)
        aligned["elevation_bin"] = _bin_series(
            aligned["6_2_elevation"], elevation_tolerance
        )
        aligned = (
            aligned.groupby(["azimuth_bin", "elevation_bin"], as_index=False)
            .agg({"5_10_signal_noise_ratio": "mean"})
            .rename(
                columns={
                    "5_10_signal_noise_ratio": f"Orbit {orbit} SNR",
                }
            )
        )
        aligned_frames.append(aligned)

    if not aligned_frames:
        return pd.DataFrame()

    merged = aligned_frames[0]
    for aligned in aligned_frames[1:]:
        merged = merged.merge(aligned, on=["azimuth_bin", "elevation_bin"], how="inner")

    return merged.sort_values(["azimuth_bin", "elevation_bin"]).reset_index(drop=True)


def _add_group_chart_sheet(
    workbook,
    source_sheet,
    sheet_name: str,
    output_frame: pd.DataFrame,
) -> None:
    if output_frame.empty:
        return

    chart_sheet = workbook.create_sheet(title=f"{sheet_name} Chart")
    ascending_chart = ScatterChart()
    descending_chart = ScatterChart()
    for chart, title in (
        (ascending_chart, f"{sheet_name} SNR vs Elevation (Ascending)"),
        (descending_chart, f"{sheet_name} SNR vs Elevation (Descending)"),
    ):
        chart.title = title
        chart.x_axis.title = "Elevation (deg)"
        chart.y_axis.title = "Signal-to-noise ratio"
        chart.x_axis.tickLblPos = "nextTo"
        chart.y_axis.tickLblPos = "nextTo"
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.x_axis.majorTickMark = "out"
        chart.y_axis.majorTickMark = "out"
        chart.x_axis.scaling.min = 5

    headers = list(output_frame.columns)
    elevation_cols: Dict[str, int] = {}
    signed_elevation_cols: Dict[str, int] = {}
    snr_cols: Dict[str, int] = {}
    for idx, name in enumerate(headers, start=1):
        if name.endswith("6_2_elevation"):
            elevation_cols[name.split(" ")[1]] = idx
        elif name.endswith("6_2_elevation_signed"):
            signed_elevation_cols[name.split(" ")[1]] = idx
        elif name.endswith("5_10_signal_noise_ratio"):
            snr_cols[name.split(" ")[1]] = idx

    data_start_row = 3
    data_end_row = 2 + output_frame.shape[0]
    elevation_values = [
        output_frame[col] for col in output_frame.columns if col.endswith("6_2_elevation")
    ]
    if elevation_values:
        max_elevation = max(
            (_max_center_elevation(series) for series in elevation_values),
            default=float("nan"),
        )
        if pd.notna(max_elevation):
            max_elevation = max(5, math.ceil(float(max_elevation)))
            ascending_chart.x_axis.scaling.max = max_elevation
            descending_chart.x_axis.scaling.max = max_elevation

    helper_start_col = len(headers) + 1
    current_col = helper_start_col
    for orbit in sorted(snr_cols.keys(), key=int):
        elev_col = elevation_cols.get(orbit)
        signed_col = signed_elevation_cols.get(orbit)
        snr_col = snr_cols.get(orbit)
        if elev_col is None or signed_col is None or snr_col is None:
            continue

        asc_elev_col = current_col
        asc_snr_col = current_col + 1
        desc_elev_col = current_col + 2
        desc_snr_col = current_col + 3
        current_col += 4

        source_sheet.cell(
            row=2, column=asc_elev_col, value=f"Orbit {orbit} elevation ascending"
        )
        source_sheet.cell(
            row=2, column=asc_snr_col, value=f"Orbit {orbit} SNR ascending"
        )
        source_sheet.cell(
            row=2, column=desc_elev_col, value=f"Orbit {orbit} elevation descending"
        )
        source_sheet.cell(
            row=2, column=desc_snr_col, value=f"Orbit {orbit} SNR descending"
        )

        elev_series = output_frame[f"Orbit {orbit} 6_2_elevation"].reset_index(drop=True)
        snr_series = output_frame[
            f"Orbit {orbit} 5_10_signal_noise_ratio"
        ].reset_index(drop=True)
        signed_series = output_frame[
            f"Orbit {orbit} 6_2_elevation_signed"
        ].reset_index(drop=True)
        signed_series = pd.to_numeric(signed_series, errors="coerce")

        for row_offset, (elev, snr, signed) in enumerate(
            zip(elev_series, snr_series, signed_series), start=data_start_row
        ):
            if pd.isna(signed):
                continue
            if signed < 0:
                source_sheet.cell(row=row_offset, column=asc_elev_col, value=elev)
                source_sheet.cell(row=row_offset, column=asc_snr_col, value=snr)
            else:
                source_sheet.cell(row=row_offset, column=desc_elev_col, value=elev)
                source_sheet.cell(row=row_offset, column=desc_snr_col, value=snr)

        asc_values = Reference(
            source_sheet,
            min_col=asc_snr_col,
            min_row=data_start_row,
            max_row=data_end_row,
        )
        asc_xvalues = Reference(
            source_sheet,
            min_col=asc_elev_col,
            min_row=data_start_row,
            max_row=data_end_row,
        )
        ascending_chart.series.append(
            Series(asc_values, asc_xvalues, title=f"Orbit {orbit}")
        )

        desc_values = Reference(
            source_sheet,
            min_col=desc_snr_col,
            min_row=data_start_row,
            max_row=data_end_row,
        )
        desc_xvalues = Reference(
            source_sheet,
            min_col=desc_elev_col,
            min_row=data_start_row,
            max_row=data_end_row,
        )
        descending_chart.series.append(
            Series(desc_values, desc_xvalues, title=f"Orbit {orbit}")
        )

    chart_sheet.add_chart(ascending_chart, "A1")
    chart_sheet.add_chart(descending_chart, "A20")


def _group_by_repeat_cycle(
    items: List[Tuple[pd.DataFrame, pd.Timestamp]],
    period: pd.Timedelta,
    tolerance: pd.Timedelta,
) -> List[List[Tuple[pd.DataFrame, pd.Timestamp]]]:
    """Group passes that fall in the same slot of the ground-track repeat cycle.

    Each pass gets a phase = (start - first start) mod *period*. Passes whose
    phases are within *tolerance* of each other (also across the cycle
    boundary) form one group, so passes a few minutes apart are no longer
    split just because they straddle an hour boundary.
    """
    if not items:
        return []
    origin = min(start for _, start in items)
    phased = sorted(
        (((start - origin) % period, frame, start) for frame, start in items),
        key=lambda entry: (entry[0], entry[2]),
    )
    groups: List[List[Tuple[pd.Timedelta, pd.DataFrame, pd.Timestamp]]] = [[phased[0]]]
    for entry in phased[1:]:
        if entry[0] - groups[-1][-1][0] <= tolerance:
            groups[-1].append(entry)
        else:
            groups.append([entry])
    if len(groups) > 1 and groups[0][0][0] + period - groups[-1][-1][0] <= tolerance:
        groups[0] = groups.pop() + groups[0]
    return [
        sorted(((frame, start) for _, frame, start in group), key=lambda item: item[1])
        for group in groups
    ]


def _frames_from_existing_output(output_path: Path) -> List[Tuple[pd.DataFrame, pd.Timestamp]]:
    """Rebuild per-orbit frames from the "Group N" sheets of a previous output.

    Only the data block is read (it ends at the first empty row, before the
    correlation table); helper and chart columns are ignored.
    """
    items: List[Tuple[pd.DataFrame, pd.Timestamp]] = []
    sheets = pd.read_excel(output_path, sheet_name=None, header=1)
    for sheet_name, df in sheets.items():
        if not re.fullmatch(r"Group \d+", sheet_name) or df.empty:
            continue
        empty_rows = df.index[df.isna().all(axis=1)]
        if len(empty_rows):
            df = df.loc[: empty_rows[0] - 1]
        orbits = sorted({
            int(m.group(1))
            for col in df.columns
            if (m := re.fullmatch(r"Orbit (\d+) time_iso_utc", str(col)))
        })
        for orbit in orbits:
            cols = {key: f"Orbit {orbit} {key}" for key in (TIME_COL, AZ_COL, EL_COL, SNR_COL)}
            if not all(c in df.columns for c in cols.values()):
                continue
            frame = df[list(cols.values())].rename(columns={v: k for k, v in cols.items()})
            frame[TIME_COL] = pd.to_datetime(frame[TIME_COL], errors="coerce", utc=True)
            for key in (AZ_COL, EL_COL, SNR_COL):
                frame[key] = pd.to_numeric(frame[key], errors="coerce")
            frame = frame.dropna().reset_index(drop=True)
            if frame.empty:
                continue
            frame["orbit"] = orbit
            items.append((frame[[TIME_COL, SNR_COL, AZ_COL, EL_COL, "orbit"]], frame[TIME_COL].min()))
    return items


def run_analysis(
    folder,
    output=None,
    azimuth_tolerance: float = 0.1,
    elevation_tolerance: float = 0.1,
    repeat_days: float = 9.0,
    group_tolerance_min: float = 30.0,
    replace: bool = False,
) -> Dict[str, object]:
    """Analyse every MEOS Extract workbook in *folder* and write *output*.

    *output* defaults to ``meos_processed.xlsx`` inside *folder*. Orbits
    already in an existing output are kept unless *replace* is true.
    Returns a summary (output path, orbits, groups, rows, skipped files).
    Raises ``ValueError`` when no usable workbook is found.
    """
    folder = Path(folder).resolve()
    if not folder.is_dir():
        raise ValueError(f"Folder not found: {folder}")

    output_path = Path(output).resolve() if output else folder / DEFAULT_OUTPUT_NAME
    excel_files = [
        *folder.glob("*.xlsx"),
        *folder.glob("*.xls"),
        *folder.glob("*.xlsm"),
    ]
    excel_files = [
        path for path in excel_files
        if path.resolve() != output_path and not path.name.startswith("~$")  # Excel lock files
    ]
    if not excel_files:
        raise ValueError(f"No Excel files found in {folder}")

    processed: Dict[int, Tuple[pd.DataFrame, pd.Timestamp]] = {}
    skipped: List[str] = []
    for path in sorted(excel_files):
        logger.info("Analysing %s", path.name)
        try:
            frame = _process_file(path)
            start_time = _extract_start_time(path)
        except ValueError as exc:
            # Other workbooks in the folder (statistics, plot index, ...)
            # are not MEOS Extract exports: skip them instead of aborting.
            logger.info("Skipped %s (not a MEOS Extract export): %s", path.name, exc)
            skipped.append(path.name)
            continue
        if frame.empty:
            logger.warning("Skipped %s: no SNR samples matched with azimuth/elevation.", path.name)
            skipped.append(path.name)
            continue
        orbit = int(frame["orbit"].iloc[0])
        if orbit in processed:
            logger.warning("Orbit %s found in more than one file; using %s.", orbit, path.name)
        processed[orbit] = (frame, start_time)

    if output_path.exists() and not replace:
        kept = 0
        for frame, start_time in _frames_from_existing_output(output_path):
            orbit = int(frame["orbit"].iloc[0])
            if orbit not in processed:
                processed[orbit] = (frame, start_time)
                kept += 1
        if kept:
            logger.info("Kept %d orbit(s) from existing %s", kept, output_path.name)

    if not processed:
        raise ValueError(f"No usable MEOS Extract files found in {folder}")

    groups = _group_by_repeat_cycle(
        list(processed.values()),
        pd.Timedelta(days=repeat_days),
        pd.Timedelta(minutes=group_tolerance_min),
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for idx, group_items in enumerate(groups):
            sheet_name = f"Group {idx + 1}"
            output_frame = None
            for frame, _ in group_items:
                orbit = int(frame["orbit"].iloc[0])
                columns = [
                    "time_iso_utc",
                    "6_1_azimuth",
                    "6_2_elevation",
                    "5_10_signal_noise_ratio",
                ]
                block = frame[columns].copy().reset_index(drop=True)
                block["time_iso_utc"] = block["time_iso_utc"].dt.tz_localize(None)
                block["6_2_elevation_signed"] = _signed_elevation(
                    block["6_2_elevation"]
                )
                block = block.rename(
                    columns={
                        "time_iso_utc": f"Orbit {orbit} time_iso_utc",
                        "6_1_azimuth": f"Orbit {orbit} 6_1_azimuth",
                        "6_2_elevation": f"Orbit {orbit} 6_2_elevation",
                        "6_2_elevation_signed": (
                            f"Orbit {orbit} 6_2_elevation_signed"
                        ),
                        "5_10_signal_noise_ratio": (
                            f"Orbit {orbit} 5_10_signal_noise_ratio"
                        ),
                    }
                )
                if output_frame is None:
                    output_frame = block
                else:
                    output_frame = pd.concat(
                        [output_frame, block], axis=1, ignore_index=False
                    )
            if output_frame is None:
                output_frame = pd.DataFrame()
            output_frame.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
                startrow=1,
            )
            worksheet = writer.sheets[sheet_name]
            if output_frame.shape[1] > 0:
                worksheet.merge_cells(
                    start_row=1,
                    start_column=1,
                    end_row=1,
                    end_column=output_frame.shape[1],
                )
                title_cell = worksheet.cell(row=1, column=1)
                title_cell.value = sheet_name
                title_cell.alignment = Alignment(horizontal="center")

            _add_group_chart_sheet(writer.book, worksheet, sheet_name, output_frame)

            frames_only = [frame for frame, _ in group_items]
            correlation_frame = _build_correlation_table(
                frames_only,
                azimuth_tolerance,
                elevation_tolerance,
            )
            if not correlation_frame.empty:
                data_end_row = 2 + output_frame.shape[0]
                title_row = data_end_row + 2
                worksheet.cell(
                    row=title_row,
                    column=1,
                    value="SNR Correlation (Azimuth/Elevation matched)",
                )
                header_row = title_row + 1
                for row_idx, row in enumerate(
                    dataframe_to_rows(correlation_frame, index=False, header=True),
                    start=header_row,
                ):
                    for col_idx, value in enumerate(row, start=1):
                        worksheet.cell(row=row_idx, column=col_idx, value=value)

                chart = ScatterChart()
                chart.title = "SNR by Orbit (Matched Az/El)"
                chart.x_axis.title = "Azimuth (binned)"
                chart.y_axis.title = "SNR"
                chart.x_axis.tickLblPos = "nextTo"
                chart.y_axis.tickLblPos = "nextTo"
                chart.x_axis.delete = False
                chart.y_axis.delete = False
                chart.x_axis.majorTickMark = "out"
                chart.y_axis.majorTickMark = "out"
                data_start_row = header_row + 1
                data_end_row = data_start_row + correlation_frame.shape[0] - 1
                azimuth_col = 1
                for col in range(3, correlation_frame.shape[1] + 1):
                    values = Reference(
                        worksheet,
                        min_col=col,
                        min_row=data_start_row,
                        max_row=data_end_row,
                    )
                    xvalues = Reference(
                        worksheet,
                        min_col=azimuth_col,
                        min_row=data_start_row,
                        max_row=data_end_row,
                    )
                    series = Series(values, xvalues, title_from_data=False)
                    label_value = worksheet.cell(row=header_row, column=col).value
                    if label_value is not None:
                        series.title = SeriesLabel(v=str(label_value))
                    chart.series.append(series)
                chart_anchor = f"A{data_end_row + 2}"
                worksheet.add_chart(chart, chart_anchor)

    total_rows = sum(len(frame) for frame, _ in processed.values())
    logger.info(
        "Saved %d orbit(s), %d rows, %d group(s) to %s", len(processed), total_rows, len(groups), output_path
    )
    if skipped:
        logger.info("Skipped %d file(s): %s", len(skipped), ", ".join(skipped))
    return {
        "output": output_path,
        "orbits": sorted(processed),
        "groups": len(groups),
        "rows": total_rows,
        "skipped": skipped,
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Multi-orbit SNR analysis of MEOS Extract Excel files: groups passes by "
            "ground-track repeat cycle and charts SNR vs elevation."
        )
    )
    parser.add_argument(
        "folder",
        nargs="?",
        default=".",
        help="Folder containing MEOS Extract Excel files.",
    )
    parser.add_argument(
        "-o",
        "--output",
        default=DEFAULT_OUTPUT_NAME,
        help="Output Excel file path.",
    )
    parser.add_argument(
        "--azimuth-tolerance",
        type=float,
        default=0.1,
        help="Tolerance (degrees) for matching azimuth values between orbits.",
    )
    parser.add_argument(
        "--elevation-tolerance",
        type=float,
        default=0.1,
        help="Tolerance (degrees) for matching elevation values between orbits.",
    )
    parser.add_argument(
        "--repeat-days",
        type=float,
        default=9.0,
        help="Ground-track repeat cycle in days used to group passes (default: 9).",
    )
    parser.add_argument(
        "--group-tolerance-min",
        type=float,
        default=30.0,
        help="Max difference (minutes) within the repeat cycle for passes of the same group.",
    )
    parser.add_argument(
        "--replace",
        action="store_true",
        help="Ignore an existing output file instead of adding its orbits to the new one.",
    )
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    try:
        run_analysis(
            args.folder,
            args.output,
            azimuth_tolerance=args.azimuth_tolerance,
            elevation_tolerance=args.elevation_tolerance,
            repeat_days=args.repeat_days,
            group_tolerance_min=args.group_tolerance_min,
            replace=args.replace,
        )
    except ValueError as exc:
        raise SystemExit(str(exc))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
